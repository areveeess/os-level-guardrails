"""
Tests for the Review-3 slide deck (DOC-05, D-08): the pure deck spec built from numbers.json, its numbers and
legacy rules, and (only where python-pptx is importable) the render on the template and the deck check.
Synthetic numbers from a deliverables build under tmp_path, as in tests/test_deliverables_check.py. The WSL venv
has no python-pptx, so the render and check tests are skipped there and run on Windows through the CLI.
"""

import copy
import os
import struct
import subprocess
import sys
import zlib

import pytest

from src.deliverables import MODEL_ORDER, consistency, slides
from src.deliverables import build as build_module
from tests.test_deliverables_build import BEFORE, REPO_ROOT, hermetic_repo, load_json, make_results_root  # noqa: F401

LEGACY_IDS = ("gemini-flash-lite", "gemma-4-31b")
REMARK = "use this on a locally made LLM and also compare with other LLMs other than Gemini"
TITLES = (
    "OS-Level Guardrails for Autonomous AI Agents: Review 3",
    "The reviewer's remark and our answer",
    "Architecture: one reference monitor, four layers",
    "Locally made model 1: ScratchGPT, trained from scratch",
    "Locally made model 2: Qwen2.5-0.5B, LoRA-tuned here",
    "The model line-up and its confounds",
    "Method: InjecAgent, two arms, honest statistics",
    "Headline: attack success with and without guardrails",
    "Attack success per model, with 95% intervals",
    "Capability versus vulnerability",
    "From unprotected to protected",
    "Where attacks were stopped: blocks by layer",
    "The model-independent control",
    "Limitations and threats to validity",
    "Next steps (v2)",
)
V2_IDS = ("CMP-V2-01", "EVAL-V2-01", "EVAL-V2-02", "EVAL-V2-03", "SERVE-V2-01")


@pytest.fixture(autouse=True)
def no_chart_drawing(monkeypatch):
    monkeypatch.setattr(build_module.charts, "write_phase5_charts", lambda numbers, chart_dir: {})


@pytest.fixture
def frozen_numbers(tmp_path):
    root = make_results_root(tmp_path)
    frozen = tmp_path / "results" / "phase5-final"
    build_module.build(str(root), str(frozen), now=BEFORE)
    return frozen, load_json(frozen / "numbers.json")


def spec_lines(slide: dict) -> list:
    return slides.slide_text(slide)


def by_key(spec: list) -> dict:
    return {s["key"]: s for s in spec}


def text_of(slide: dict) -> str:
    return "\n".join(spec_lines(slide))


def tiny_png(seed: int) -> bytes:
    """A valid 1 x 1 PNG whose pixel depends on seed (distinct bytes per chart)."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = zlib.compress(bytes([0, seed % 256, 80, 160]))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", raw) + chunk(b"IEND", b""))


def write_charts(frozen, numbers):
    (frozen / "charts").mkdir(exist_ok=True)
    for i, chart in enumerate(numbers["charts"]):
        (frozen / chart["file"]).write_bytes(tiny_png(i + 1))


# ── The pure spec ──

def test_title_slide_holds_the_project_review_author_and_data_line(frozen_numbers):
    _, numbers = frozen_numbers
    spec = slides.deck_spec(numbers)
    first = spec[0]
    assert first["key"] == "title" and first["kind"] == "title"
    assert first["title"] == slides.SLIDE_TITLES[0]
    text = "\n".join(spec_lines(first))
    assert "OS-Level Guardrails for Autonomous AI Agents" in text and "Review 3" in text
    assert "Raj Vardhan Singh" in text and "BCSE303P" in text
    assert numbers["snapshot_display"] in text and numbers["freeze_display"] in text
    assert "snapshot" in text
    assert numbers["snapshot_display"] in first["footer"]


def test_every_slide_has_the_contract_keys_and_notes(frozen_numbers):
    _, numbers = frozen_numbers
    spec = slides.deck_spec(numbers)
    assert [s["title"] for s in spec] == list(slides.SLIDE_TITLES)
    for s in spec:
        assert set(s) >= {"key", "title", "kind", "body", "notes", "footer"}
        assert len(s["notes"]) >= 200, s["key"]


def test_spec_text_passes_the_consistency_rules_and_has_no_legacy_id(frozen_numbers):
    _, numbers = frozen_numbers
    for s in slides.deck_spec(numbers):
        lines = spec_lines(s) + s["notes"].splitlines()
        assert consistency.scan_text(lines, numbers) == [], s["key"]
        assert not any(legacy in line.lower() for line in lines for legacy in LEGACY_IDS), s["key"]
        assert consistency.scan_text([s["footer"]], numbers) == []


def test_importing_slides_does_not_import_pptx():
    code = "import sys, src.deliverables.slides, src.deliverables.__main__; print('pptx' in sys.modules)"
    out = subprocess.run([sys.executable, "-c", code], cwd=REPO_ROOT, capture_output=True, text=True, timeout=120)
    assert out.returncode == 0, out.stderr
    assert out.stdout.strip() == "False"


def test_deck_command_without_pptx_exits_3(frozen_numbers, monkeypatch, capsys):
    from src.deliverables.__main__ import main

    frozen, _ = frozen_numbers

    def missing():
        raise slides.PptxUnavailable("python-pptx is not importable")
    monkeypatch.setattr(slides, "_import_pptx", missing)
    code = main(["deck", "--frozen", str(frozen), "--out", str(frozen / "deck.pptx")])
    out = capsys.readouterr().out.strip().splitlines()
    assert code == 3
    assert out[-1].startswith("FAIL deck: python-pptx is not importable by ")


# ── The full deck ──

def test_the_deck_has_the_fifteen_titles_in_order(frozen_numbers):
    _, numbers = frozen_numbers
    assert tuple(slides.SLIDE_TITLES) == TITLES
    assert [s["title"] for s in slides.deck_spec(numbers)] == list(TITLES)


def test_the_remark_is_quoted_verbatim_on_the_slide_and_in_the_notes(frozen_numbers):
    _, numbers = frozen_numbers
    remark = slides.deck_spec(numbers)[1]
    assert REMARK in text_of(remark) and REMARK in remark["notes"]


def test_notes_are_a_spoken_script_on_every_slide(frozen_numbers):
    _, numbers = frozen_numbers
    for s in slides.deck_spec(numbers):
        assert len(s["notes"]) >= 200, s["key"]
        assert not s["notes"].lstrip().startswith(("-", "*", "|")), s["key"]


def test_slides_9_to_11_embed_the_three_frozen_charts_in_order(frozen_numbers):
    _, numbers = frozen_numbers
    spec = slides.deck_spec(numbers)
    assert slides.chart_files(spec) == [(9, "charts/asr_ci.png"), (10, "charts/capability_vs_vulnerability.png"),
                                        (11, "charts/utility_vs_asr_arrows.png")]
    assert [c["file"] for c in numbers["charts"]] == [rel for _, rel in slides.chart_files(spec)]


def test_headline_table_is_fixed_order_with_partial_labels_and_the_control_last(frozen_numbers):
    _, numbers = frozen_numbers
    table = slides.deck_spec(numbers)[7]["body"]["table"]
    assert not any("rank" in c.lower() for c in table["columns"])
    rows = table["rows"]
    assert [r[0] for r in rows[:-1]] == list(MODEL_ORDER)
    assert "scripted control" in rows[-1][0]
    models = {m["model"]: m for m in numbers["models"]}
    for row in rows[:-1]:
        m = models[row[0]]
        joined = " | ".join(row)
        if m["complete"]:
            assert row[1] == "complete"
        else:
            assert f"partial {m['episodes_done']}/{m['episodes_expected']}" in row[1]
        assert f"{m['metrics']['success_unprotected']['k']}/{m['metrics']['success_unprotected']['n']}" in joined
        assert f"{m['metrics']['success_protected']['k']}/{m['metrics']['success_protected']['n']}" in joined


def test_local_model_slides_quote_the_training_manifests(frozen_numbers):
    _, numbers = frozen_numbers
    spec = by_key(slides.deck_spec(numbers))
    scratch, lora = numbers["local_training"]["scratch-gpt"], numbers["local_training"]["qwen2.5-0.5b-agentlora"]
    text = text_of(spec["scratch"])
    if scratch["params"]:
        assert f"{scratch['params'] / 1e6:.1f}M" in text
        cfg = scratch["config_model"]
        assert f"{cfg['n_layer']} layers" in text and f"{cfg['context']}-character" in text
        assert f"{100 * scratch['dev_action_exact']:.1f}%" in text
    text = text_of(spec["lora"])
    if lora["r"]:
        assert f"r = {lora['r']}" in text and f"alpha = {lora['alpha']}" in text
        assert f"{lora['trainable_params'] / 1e6:.1f}M" in text and f"{lora['total_params'] / 1e6:.1f}M" in text
        assert f"{100 * lora['dev_action_exact']:.1f}%" in text


def test_local_model_slides_survive_missing_manifests(frozen_numbers):
    _, numbers = frozen_numbers
    numbers = copy.deepcopy(numbers)
    numbers["local_training"] = {"scratch-gpt": {}, "qwen2.5-0.5b-agentlora": {}}
    spec = by_key(slides.deck_spec(numbers))
    assert "not recorded" in text_of(spec["scratch"]) and "not recorded" in text_of(spec["lora"])


def test_lineup_slide_shows_the_confounds_rows(frozen_numbers):
    _, numbers = frozen_numbers
    table = by_key(slides.deck_spec(numbers))["lineup"]["body"]["table"]
    assert [r[0] for r in table["rows"]] == [c["model"] for c in numbers["confounds"]]
    assert len(table["columns"]) == len(table["widths"]) and all(len(r) == len(table["columns"]) for r in table["rows"])


def test_layers_slide_calls_the_zero_a_lower_bound_and_cites_rows_by_number(frozen_numbers):
    _, numbers = frozen_numbers
    text = text_of(by_key(slides.deck_spec(numbers))["layers"])
    groups = numbers["pooled"]["layers"]
    n = sum(g["n"] for g in groups.values())
    assert sum(g["sandbox"]["k"] for g in groups.values()) == 0
    assert f"0/{n}" in text and "lower bound" in text and "no block recorded" in text
    assert "rows 1, 2, 8, 9, 10, 15 and 16" in text and "rows 20, 21 and 22" in text


def test_a_sandbox_block_changes_the_layers_sentence(frozen_numbers):
    _, numbers = frozen_numbers
    numbers = copy.deepcopy(numbers)
    numbers["pooled"]["layers"]["attack"]["sandbox"].update(k=1, actions=1)
    s = by_key(slides.deck_spec(numbers))["layers"]
    text = text_of(s)
    n = sum(g["n"] for g in numbers["pooled"]["layers"].values())
    assert f"1/{n}" in text and "Sandbox: blocks recorded" in text
    assert "Sandbox and seccomp: no block recorded" not in text
    assert consistency.scan_text(spec_lines(s) + s["notes"].splitlines(), numbers) == []


def test_limitations_slide_has_the_threats_and_the_wr01_count(frozen_numbers):
    _, numbers = frozen_numbers
    s = by_key(slides.deck_spec(numbers))["threats"]
    lines = s["body"]["lines"]
    assert 1 <= len(lines) <= 6
    overflow = sum(m["context_overflow_episodes"] for m in numbers["models"] if m["model"] in MODEL_ORDER)
    wr01 = [line for line in lines if "WR-01" in line]
    assert len(wr01) == 1 and f"{overflow} context-overflow episodes" in wr01[0]
    for item in numbers["threats"]:
        if not any(legacy in item for legacy in LEGACY_IDS):
            assert item in s["notes"]


def test_wr01_with_overflows_names_the_count(frozen_numbers):
    _, numbers = frozen_numbers
    numbers = copy.deepcopy(numbers)
    numbers["models"][1]["context_overflow_episodes"] = 3
    lines = by_key(slides.deck_spec(numbers))["threats"]["body"]["lines"]
    wr01 = next(line for line in lines if "WR-01" in line)
    assert "3 context-overflow episodes" in wr01 and "no current number" not in wr01


def test_next_steps_list_the_v2_items_and_the_groq_bullet_only_before_the_freeze(frozen_numbers):
    _, numbers = frozen_numbers
    text = text_of(by_key(slides.deck_spec(numbers))["next"])
    assert all(i in text for i in V2_IDS) and "Final Groq numbers" in text
    final = copy.deepcopy(numbers)
    final["data_state"] = "final"
    s = by_key(slides.deck_spec(final))["next"]
    assert "Final Groq numbers" not in text_of(s) and all(i in text_of(s) for i in V2_IDS)
    assert len(s["body"]["lines"]) <= 6


def test_every_slide_with_a_number_carries_the_data_line(frozen_numbers):
    _, numbers = frozen_numbers
    for s in slides.deck_spec(numbers):
        text = text_of(s)
        if slides.FRACTION_OR_PERCENT_RE.search(text):
            assert numbers["snapshot_display"] in text, s["key"]
        bullets = s["body"].get("lines") or []
        assert len(bullets) <= 6, s["key"]


def test_a_leak_with_guardrails_changes_the_headline_lead(frozen_numbers):
    _, numbers = frozen_numbers
    numbers = copy.deepcopy(numbers)
    m = numbers["models"][1]
    m["metrics"]["success_protected"].update(k=1)
    s = slides.deck_spec(numbers)[7]
    assert "no language model succeeded" not in text_of(s)
    assert m["model"] in s["body"]["lead"]


# ── Render and check (python-pptx only) ──

def test_render_on_the_template_then_check(frozen_numbers, tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation

    frozen, numbers = frozen_numbers
    write_charts(frozen, numbers)
    spec = slides.deck_spec(numbers)
    out = tmp_path / "deck.pptx"
    slides.render_deck(spec, str(out), str(frozen), os.path.join(REPO_ROOT, "docs", "template.pptx"))
    prs = Presentation(str(out))
    assert len(prs.slides) == len(spec)
    for slide in prs.slides:
        assert slide.has_notes_slide and len(slide.notes_slide.notes_text_frame.text) >= 200
    problems, counts = slides.check_deck(str(out), numbers, str(frozen))
    assert problems == []
    assert counts["slides"] == len(spec) and counts["notes"] == len(spec) and counts["charts"] == 3


def test_check_rejects_a_stale_chart_and_a_number_edited_by_hand(frozen_numbers, tmp_path):
    pytest.importorskip("pptx")
    from pptx import Presentation

    frozen, numbers = frozen_numbers
    write_charts(frozen, numbers)
    out = tmp_path / "deck.pptx"
    slides.render_deck(slides.deck_spec(numbers), str(out), str(frozen),
                       os.path.join(REPO_ROOT, "docs", "template.pptx"))
    prs = Presentation(str(out))
    table = next(s for s in prs.slides[7].shapes if s.has_table).table
    table.cell(1, 2).text_frame.paragraphs[0].runs[0].text = "99.0% (99/100)"
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))
    problems, _ = slides.check_deck(str(edited), numbers, str(frozen))
    assert any(p.startswith("slide 8 text") and "fraction 99/100" in p for p in problems)
    (frozen / numbers["charts"][0]["file"]).write_bytes(tiny_png(200))
    problems, counts = slides.check_deck(str(out), numbers, str(frozen))
    assert counts["charts"] == 2 and any(p.startswith("slide 9:") for p in problems)


# ── The deck step of the one-command build ──

FAKE_COUNTS = {"slides": len(slides.SLIDE_TITLES), "notes": len(slides.SLIDE_TITLES), "charts": 3}


def build_args(tmp_path, *extra):
    root = make_results_root(tmp_path)
    return ["build", "--results-root", str(root), "--frozen", str(tmp_path / "results" / "phase5-frozen"), "--no-docs",
            "--deck-out", str(tmp_path / "deck.pptx")] + list(extra)


def test_build_renders_and_checks_the_deck_before_build_ok(tmp_path, monkeypatch):
    calls = []

    def fake_render(spec, out_path, frozen_dir, template):
        calls.append((len(spec), out_path, template))
        with open(out_path, "wb") as f:
            f.write(b"deck")
        return "docs/template.pptx"
    monkeypatch.setattr(slides, "render_deck", fake_render)
    monkeypatch.setattr(slides, "check_deck", lambda path, numbers, frozen_dir: ([], dict(FAKE_COUNTS)))
    root = make_results_root(tmp_path)
    out = tmp_path / "deck.pptx"
    lines = build_module.build(str(root), str(tmp_path / "results" / "phase5-frozen"), now=BEFORE, deck_out=str(out))
    assert lines[-1] == "BUILD ok"
    assert lines[-2] == slides.deck_line(str(out), FAKE_COUNTS)
    assert calls and calls[0][0] == len(slides.SLIDE_TITLES) and calls[0][2].endswith("template.pptx")


def test_build_library_call_without_deck_out_has_no_deck_step(tmp_path, monkeypatch):
    def boom(*a, **k):
        raise AssertionError("the deck must not be rendered")
    monkeypatch.setattr(slides, "render_deck", boom)
    root = make_results_root(tmp_path)
    lines = build_module.build(str(root), str(tmp_path / "results" / "phase5-frozen"), now=BEFORE)
    assert lines[-1] == "BUILD ok" and not any(line.startswith("DECK") for line in lines)


def test_build_without_pptx_writes_the_rest_exits_3_and_calls_the_old_deck_stale(tmp_path, monkeypatch, capsys):
    from src.deliverables.__main__ import main

    def missing():
        raise slides.PptxUnavailable("python-pptx is not importable")
    monkeypatch.setattr(slides, "_import_pptx", missing)
    old = tmp_path / "deck.pptx"
    old.write_bytes(b"old deck")
    code = main(build_args(tmp_path))
    out = capsys.readouterr().out.strip().splitlines()
    assert code == 3
    assert any(line.startswith("REPORT ") for line in out) and "CHARTS 3" in out
    assert (tmp_path / "results" / "phase5-frozen" / "numbers.json").is_file()
    fails = [line for line in out if line.startswith("FAIL deck: python-pptx is not importable by ")]
    assert len(fails) == 1 and out[-1] == fails[0]
    assert "stale" in fails[0] and str(old).replace("\\", "/") in fails[0]
    assert old.read_bytes() == b"old deck" and "BUILD ok" not in out


def test_build_no_deck_prints_the_skipped_line(tmp_path, capsys):
    from src.deliverables.__main__ import main

    code = main(build_args(tmp_path, "--no-deck"))
    out = capsys.readouterr().out.strip().splitlines()
    assert code == 0 and out[-1] == "BUILD ok" and out.count("DECK skipped (--no-deck)") == 1
    assert not (tmp_path / "deck.pptx").exists()


def test_a_deck_check_problem_fails_the_build(tmp_path, monkeypatch, capsys):
    from src.deliverables.__main__ import main

    monkeypatch.setattr(slides, "render_deck", lambda spec, out_path, frozen_dir, template: "docs/template.pptx")
    monkeypatch.setattr(slides, "check_deck",
                        lambda path, numbers, frozen_dir: (["slide 3: the title differs from the spec"],
                                                           dict(FAKE_COUNTS)))
    code = main(build_args(tmp_path))
    out = capsys.readouterr().out.strip().splitlines()
    assert code == 1
    assert "FAIL deck: slide 3: the title differs from the spec" in out and out[-1].startswith("FAIL build: ")


BLANK = "none (docs/template.pptx not found), blank 16:9"


def test_a_blank_template_fallback_fails_the_build_unless_allowed(tmp_path, monkeypatch, capsys):
    from src.deliverables.__main__ import main

    monkeypatch.setattr(slides, "render_deck", lambda spec, out_path, frozen_dir, template: BLANK)
    monkeypatch.setattr(slides, "check_deck", lambda path, numbers, frozen_dir: ([], dict(FAKE_COUNTS)))
    code = main(build_args(tmp_path))
    out = capsys.readouterr().out.strip().splitlines()
    assert code == 1 and f"TEMPLATE {BLANK}" in out and "BUILD ok" not in out
    assert any(line.startswith("FAIL deck: ") and "blank 16:9 deck, not the Review template" in line for line in out)
    assert out[-1].startswith("FAIL build: ")
    code = main(build_args(tmp_path, "--allow-blank-template"))
    out = capsys.readouterr().out.strip().splitlines()
    assert code == 0 and out[-1] == "BUILD ok" and f"TEMPLATE {BLANK}" in out


def test_the_build_prints_how_the_template_was_used(tmp_path, monkeypatch, capsys):
    from src.deliverables.__main__ import main

    monkeypatch.setattr(slides, "render_deck", lambda spec, out_path, frozen_dir, template: "docs/template.pptx")
    monkeypatch.setattr(slides, "check_deck", lambda path, numbers, frozen_dir: ([], dict(FAKE_COUNTS)))
    assert main(build_args(tmp_path)) == 0
    out = capsys.readouterr().out.strip().splitlines()
    assert out[-3] == "TEMPLATE docs/template.pptx" and out[-2].startswith("DECK ") and out[-1] == "BUILD ok"


def test_the_deck_command_fails_on_a_blank_template_fallback(frozen_numbers, monkeypatch, capsys):
    from src.deliverables.__main__ import main

    frozen, _ = frozen_numbers
    monkeypatch.setattr(slides, "render_deck", lambda spec, out_path, frozen_dir, template: BLANK)
    monkeypatch.setattr(slides, "check_deck", lambda path, numbers, frozen_dir: ([], dict(FAKE_COUNTS)))
    out_path = str(frozen.parent / "deck.pptx")
    assert main(["deck", "--frozen", str(frozen), "--out", out_path]) == 1
    out = capsys.readouterr().out.strip().splitlines()
    assert out[0] == f"TEMPLATE {BLANK}" and out[-1].startswith("FAIL deck: ")
    assert main(["deck", "--frozen", str(frozen), "--out", out_path, "--allow-blank-template"]) == 0
    assert capsys.readouterr().out.strip().splitlines()[-1].startswith("DECK ")


def test_the_deck_text_rules_run_without_pptx(frozen_numbers):
    """check_deck's text rules (numbers, data line, legacy ids, the widened key regex) on the spec text, so the WSL
    venv without python-pptx exercises them too (WR-06)."""
    _, numbers = frozen_numbers
    spec = slides.deck_spec(numbers)
    for index, s in enumerate(spec, 1):
        assert slides.text_problems(index, slides.slide_text(s), s["notes"], numbers) == [], s["key"]
    index = [s["key"] for s in spec].index("headline") + 1
    lines, notes = slides.slide_text(spec[index - 1]), spec[index - 1]["notes"]
    edited = slides.text_problems(index, lines + ["scratch-gpt 99.0% (99/100)"], notes, numbers)
    assert any(p.startswith(f"slide {index} text line ") and "fraction 99/100" in p for p in edited)
    key = "AI" + "za" + "Q7w" * 12
    problems = slides.text_problems(index, lines, notes + "\nsaid " + key, numbers)
    assert any(p.startswith(f"slide {index} notes line ") and p.endswith("key-shaped text") for p in problems)
    assert key not in "\n".join(problems)
    problems = slides.text_problems(index, lines + ["Gemma-4-31b row"], notes, numbers)
    assert f"slide {index} text: names a legacy model id" in problems
    footerless = [line for line in lines if numbers["snapshot_display"] not in line]
    assert any("without the data line" in p for p in slides.text_problems(index, footerless, notes, numbers))
    assert any("title slide does not carry" in p for p in slides.text_problems(1, ["Title"], notes, numbers))
