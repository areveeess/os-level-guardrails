"""
Deliverables slides — the Review-3 deck generated from numbers.json and the frozen charts
========================================================================================
The Review-3 slide deck (DOC-05) quotes the same results as the report, the README, the SDD
and the viva pages, so it is generated the same way (D-03, D-08): deck_spec turns numbers.json
into a list of plain slide dicts, render_deck draws them with python-pptx on the Review-2
template, and check_deck reopens the written file and fails on any number that is not in
numbers.json, any chart that is not byte-identical to a frozen chart, any legacy model id, any
slide without speaker notes, and any key-shaped text. A number typed onto a slide by hand would
drift from the data while the Groq loop runs until the D-09 freeze; a generated deck cannot.

deck_spec is pure and never imports python-pptx, so the spec can be tested in the WSL venv (which
has no python-pptx). render_deck and check_deck import it inside the function; when it is missing
they raise PptxUnavailable, which the CLI and the build turn into a loud exit 3 rather than
silently leaving an old deck in place.

Design: the template's layouts carry gradient-filled decoration on a dark background, so every
slide uses the template's plain BLANK layout with the chart surface colour as background, its
own title, one thin accent rule and a footer with the data line (D-01). The deck file is a
binary output and is never committed.
"""

import hashlib
import math
import os
import re

from src.benchmark.report import GRID, LAYER_EVIDENCE, LAYER_TITLES, LAYERS, SURFACE, TEXT_PRIMARY, TEXT_SECONDARY
from src.deliverables import CONTROL_MODEL, MODEL_ORDER
from src.deliverables.blocks import data_line
from src.deliverables.consistency import scan_text

DEFAULT_OUT = os.path.join("docs", "OS_Level_Guardrails_Review_3.pptx")
DEFAULT_TEMPLATE = os.path.join("docs", "template.pptx")

SLIDE_TITLES = (
    "OS-Level Guardrails for Autonomous AI Agents: Review 3",
    "The reviewer's remark and our answer",
    "Architecture: one reference monitor, four layers",
    "Locally made model 1: ScratchGPT, trained from scratch",
    "Locally made model 2: Qwen2.5-0.5B, LoRA-tuned here",
    "The model line-up and its confounds",
    "Method: InjecAgent, two arms, honest statistics",
    "Headline: attack success with and without guardrails",
    "Attack success per model, with 95% intervals",
    "Capability versus vulnerability",
    "From unprotected to protected",
    "Where attacks were stopped: blocks by layer",
    "The model-independent control",
    "Limitations and threats to validity",
    "Next steps (v2)",
)

AUTHOR = "Raj Vardhan Singh (24BDS0200)"
COURSE = "BCSE303P Operating Systems Lab"
# The Review-2 panel's remark, verbatim from .planning/PROJECT.md
REMARK = "use this on a locally made LLM and also compare with other LLMs other than Gemini"
LEGACY_IDS = ("gemini-flash-lite", "gemma-4-31b")
MIN_NOTES = 200
NOT_RECORDED = "not recorded"

# The v2 requirements of .planning/REQUIREMENTS.md, in plain words
V2_STEPS = (
    ("CMP-V2-01", "full benchmark runs of phi3.5 (3.8B) and gemma3 (1B) through Ollama"),
    ("EVAL-V2-01", "utility under attack on a larger task set, with case-clustered bootstrap intervals"),
    ("EVAL-V2-02", "adaptive attacks that retry after a block"),
    ("EVAL-V2-03", "replication on AgentDojo or Agent Security Bench"),
    ("SERVE-V2-01", "export the LoRA-tuned Qwen to GGUF and serve it through Ollama"),
)

# Visual system: the chart surface as the slide background, the chart text colours, one sans-serif family and
# one accent colour, the chart palette's guardrails blue (charts.BLUE, Okabe-Ito).
BG, INK, MUTED, RULE = (c.lstrip("#").upper() for c in (SURFACE, TEXT_PRIMARY, TEXT_SECONDARY, GRID))
ACCENT = "0072B2"
TINT = "EAF2F8"
DARK = "2B2B2B"
WHITE = "FFFFFF"
FONT = "Arial"
NO_TABLE_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"      # PowerPoint's "No Style, No Grid"

# Slide geometry in inches (16:9, 10 x 5.625 in, the template's size)
SLIDE_W, SLIDE_H = 10.0, 5.625
MARGIN = 0.5
CONTENT_W = SLIDE_W - 2 * MARGIN
TITLE_Y, TITLE_H = 0.3, 0.62
RULE_Y, RULE_W, RULE_H = 1.0, 0.7, 0.045
BODY_Y = 1.22
BODY_BOTTOM = 5.12
FOOTER_Y, FOOTER_H = 5.27, 0.24
CHART_X, CHART_Y, CHART_W = 0.4, 1.17, 7.0
SIDE_X, SIDE_W = 7.65, 1.95

FRACTION_OR_PERCENT_RE = re.compile(r"\d/\d|%")


class PptxUnavailable(Exception):
    """python-pptx (or its lxml) cannot be imported by this interpreter."""


# ── Values from numbers.json ──

def _llm(numbers: dict) -> list:
    return [m for m in numbers["models"] if m["model"] in MODEL_ORDER]


def _model(numbers: dict, name: str):
    return next((m for m in numbers["models"] if m["model"] == name), None)


def _kn(metric: dict) -> str:
    return f"{metric['k']}/{metric['n']}"


def _fp(k: int, n: int) -> str:
    """'92.5% (111/120)': the percentage the consistency check recomputes from k/n."""
    return "— (0/0)" if not n else f"{100.0 * k / n:.1f}% ({k}/{n})"


def _pk(metric: dict) -> str:
    return _fp(metric["k"], metric["n"])


def _run_label(m: dict) -> str:
    if m["complete"]:
        return "complete"
    label = f"partial {m['episodes_done']}/{m['episodes_expected']}"
    return label + ", error-dominated" if m["status"] == "error-dominated" else label


def _clean(numbers: dict) -> bool:
    """No language model succeeded on any attack episode with guardrails."""
    return all(m["metrics"]["success_protected"]["k"] == 0 for m in _llm(numbers))


def _leaks(numbers: dict) -> list:
    return [m for m in _llm(numbers) if m["metrics"]["success_protected"]["k"]]


def _pooled_layer(numbers: dict, layer: str) -> tuple:
    groups = numbers["pooled"]["layers"]
    return sum(g[layer]["k"] for g in groups.values()), sum(g["n"] for g in groups.values())


def _bench(numbers: dict) -> dict:
    """Cases per kind and runs per case, read back from the scripted control (it runs every case)."""
    c = numbers.get("control") or {}
    attack_cases = (c.get("cases") or {}).get("attack") or 0
    episodes = (c.get("protected") or {}).get("episodes") or 0
    runs = max(1, episodes // attack_cases) if attack_cases else 1
    kinds = c.get("by_kind") or {}
    control = _model(numbers, CONTROL_MODEL)
    benign_episodes = control["metrics"]["benign_task_unprotected"]["n"] if control else 0
    out = {"runs": runs, "dh": (kinds.get("dh") or {}).get("episodes", 0) // runs,
           "ds": (kinds.get("ds") or {}).get("episodes", 0) // runs, "benign": benign_episodes // runs}
    out["total"] = out["dh"] + out["ds"] + out["benign"]
    return out


def _millions(value) -> str:
    return f"{value / 1e6:.1f}M" if value else NOT_RECORDED


def _pct(value) -> str:
    return f"{100 * value:.1f}%" if value is not None else NOT_RECORDED


def _phase1_rows() -> dict:
    """The Phase 1 residual-bypass rows by number, read from report.LAYER_EVIDENCE (cited by number only)."""
    found = {}
    for key, pattern in (("seccomp", r"rows ([\d, and-]+?) were recorded as blocked by seccomp"),
                         ("sandbox", r"rows ([\d, and-]+?) as blocked by the sandbox"),
                         ("contained", r"in rows ([\d, and-]+?) the layer contained")):
        match = re.search(pattern, LAYER_EVIDENCE)
        if match:
            found[key] = match.group(1)
    return found


# ── The spec (pure) ──

def footer_line(numbers: dict) -> str:
    """The compact data line on every slide (D-01): the snapshot time and the data state."""
    if numbers["data_state"] == "snapshot":
        return (f"Data: snapshot of {numbers['snapshot_display']} (UTC), not final; the D-09 freeze is "
                f"{numbers['freeze_display']}. Partial rows show their episodes done.")
    return (f"Data: final, copied at {numbers['snapshot_display']} (UTC) after the D-09 freeze of "
            f"{numbers['freeze_display']}.")


def _state_sentence(numbers: dict) -> str:
    if numbers["data_state"] == "snapshot":
        return (f"The hosted runs were still going at that moment, so their rows are labelled partial and are not "
                f"final; the final freeze is at {numbers['freeze_display']} UTC.")
    return f"This is the final data, copied after the freeze at {numbers['freeze_display']} UTC."


def _slide(key: str, index: int, kind: str, body: dict, notes, numbers: dict) -> dict:
    return {"key": key, "title": SLIDE_TITLES[index], "kind": kind, "body": body,
            "notes": "\n".join(notes) if isinstance(notes, (list, tuple)) else notes,
            "footer": footer_line(numbers)}


def _title_slide(numbers: dict) -> dict:
    notes = [f"Good morning. I am Raj Vardhan Singh, and this is Review 3 of my Operating Systems lab project, "
             f"OS-level guardrails for autonomous AI agents. In Review 2 I showed a reference monitor built from a "
             f"tool broker, a policy engine, Linux namespaces and a seccomp-BPF filter, driven by Gemini. Today I "
             f"answer the panel's remark from that review: I made language models locally and compared them with "
             f"other, non-Gemini models under exactly the same guardrails.",
             f"Every number in this deck is generated by one command from a frozen snapshot of the benchmark data, "
             f"taken at {numbers['snapshot_display']} UTC. {_state_sentence(numbers)}"]
    return {
        "key": "title",
        "title": SLIDE_TITLES[0],
        "kind": "title",
        "body": {"kicker": f"{COURSE}  ·  Project Review 3",
                 "subtitle": "The same guardrails on locally made and hosted language models",
                 "byline": AUTHOR},
        "notes": "\n".join(notes),
        "footer": data_line(numbers),
    }


def _remark_slide(numbers: dict) -> dict:
    llm = _llm(numbers)
    local = [m for m in llm if m["group"] == "local"]
    hosted = [m for m in llm if m["group"] == "hosted"]
    if _clean(numbers):
        outcome = f"With guardrails, no attack succeeded for any of the {len(llm)} models, the weakest included"
        spoken = ("The short answer is the last line: with the guardrails on, not one attack succeeded for any of "
                  "these models, including the small, unaligned one I trained myself.")
    else:
        leaks = _leaks(numbers)
        outcome = ("With guardrails, attacks still succeeded for "
                   + ", ".join(f"{m['model']} {_kn(m['metrics']['success_protected'])}" for m in leaks))
        spoken = ("The last line is the honest summary: with the guardrails on, attacks still got through for "
                  f"{len(leaks)} of the models, and I show exactly where on the results slides.")
    lines = ["Two models made here: ScratchGPT, trained from scratch, and Qwen2.5-0.5B, LoRA-tuned here",
             f"{len(llm)} language models compared, none of them Gemini: {len(local)} run locally, "
             f"{len(hosted)} hosted on Groq",
             "The same system prompt, cases, seeds and arms for every model",
             outcome]
    notes = [f"In Review 2 the panel said, and I quote: \"{REMARK}\". That remark has two parts, and I answered both.",
             "First, a locally made model. I made two: ScratchGPT is a small transformer that I wrote and trained "
             "from scratch, and Qwen2.5-0.5B is a pretrained model that I fine-tuned here with LoRA.",
             f"Second, models other than Gemini. The comparison now has {len(llm)} language models, {len(local)} "
             f"running on my own laptop and {len(hosted)} hosted on Groq, and none of them is Gemini. Every model "
             f"gets the same system prompt, the same cases, the same seeds and the same two arms, so the only thing "
             f"that changes between rows is the model.",
             spoken]
    body = {"quote": f"“{REMARK}”", "attribution": "Review 2 panel", "label": "Our answer", "lines": lines}
    return _slide("remark", 1, "quote", body, notes, numbers)


def _architecture_slide(numbers: dict) -> dict:
    boxes = [
        {"role": "model", "label": "Language model",
         "detail": "untrusted: it only writes text, and the harness parses one action from it"},
        {"role": "layer", "label": "Tool broker", "detail": "only the tools this task may call"},
        {"role": "layer", "label": "Policy engine",
         "detail": "YAML rules allow, flag or block shell, file and web actions"},
        {"role": "layer", "label": "Namespaces", "detail": "user, mount, PID and network; a private filesystem"},
        {"role": "layer", "label": "seccomp-BPF", "detail": "kernel filter that refuses dangerous system calls"},
        {"role": "kernel", "label": "Linux kernel", "detail": ""},
        {"role": "log", "label": "Decision log",
         "detail": "every allow, flag and block as one JSON line; the live dashboard reads it"},
        {"role": "arms", "label": "Two arms",
         "detail": "Unprotected: no layers, an emulated OS. Protected: all four layers, the real sandbox."},
    ]
    notes = ["This is the system from Review 2, unchanged, and that is the point: the guardrails do not know or care "
             "which model is driving the agent.",
             "The language model is treated as an untrusted principal. It only produces text, and the harness parses "
             "that text into one action. Every action then meets the reference monitor. The tool broker lets a task "
             "call only the tools on its capability list. The policy engine checks every shell, file and web action "
             "against YAML rules and answers allow, flag or block. Anything that is allowed runs inside Linux user, "
             "mount, PID and network namespaces with a private filesystem, under a seccomp-BPF filter that refuses "
             "dangerous system calls.",
             "Every decision is written as one JSON line to a log that the dashboard shows live. In the benchmark, the "
             "unprotected arm runs the same agent with none of these layers against an emulated OS, which gives the "
             "baseline that the protected arm is compared with."]
    body = {"diagram": {"container": "Reference monitor: every action is checked first",
                        "boxes": boxes}}
    return _slide("architecture", 2, "diagram", body, notes, numbers)


def _stop_text(reason) -> str:
    if not reason:
        return f"stop reason {NOT_RECORDED}"
    return {"early_stop": "early stopping"}.get(reason, str(reason).replace("_", " "))


def _scratch_slide(numbers: dict) -> dict:
    t = (numbers.get("local_training") or {}).get("scratch-gpt") or {}
    cfg = t.get("config_model") or {}
    if all(cfg.get(k) for k in ("n_layer", "n_head", "d_model")):
        arch = (f"A decoder-only transformer written here in PyTorch: {cfg['n_layer']} layers, {cfg['n_head']} heads, "
                f"width {cfg['d_model']}")
    else:
        arch = f"A decoder-only transformer written here in PyTorch (architecture {NOT_RECORDED})"
    if cfg.get("vocab_size") and cfg.get("context"):
        chars = f"Character-level: a {cfg['vocab_size']}-symbol vocabulary and a {cfg['context']}-character context"
    else:
        chars = f"Character-level (vocabulary and context {NOT_RECORDED})"
    if t.get("best_step") is not None and t.get("steps") is not None and t.get("dev_action_exact") is not None:
        best = (f"Best checkpoint at step {t['best_step']} of {t['steps']} ({_stop_text(t.get('stop_reason'))}); "
                f"exact next action on the dev set: {_pct(t['dev_action_exact'])}")
    else:
        best = f"Best checkpoint and dev-set accuracy {NOT_RECORDED}"
    lines = [arch, chars,
             "Trained from scratch on synthetic agent conversations; the benchmark's tools are held out",
             best,
             "No pretraining and no safety tuning: a deliberately naive, unaligned agent"]
    params = _millions(t.get("params"))
    spoken_params = f"about {t['params'] / 1e6:.1f} million" if t.get("params") else "a few million"
    accuracy = (f"On held-out dev conversations it picks exactly the right next action {_pct(t['dev_action_exact'])} "
                f"of the time, so it really can act as an agent, if clumsily."
                if t.get("dev_action_exact") is not None else "Its dev-set accuracy was not recorded.")
    notes = [f"The first locally made model is ScratchGPT. I wrote this transformer myself in PyTorch and trained it "
             f"from random weights here, on my laptop. It has {spoken_params} parameters and reads text one "
             f"character at a time.",
             "It never saw any pretraining or safety data, so it has no idea that it should refuse anything. That "
             "makes it the hardest test for the guardrails: nothing inside the model helps them.",
             accuracy,
             "The benchmark's tools were held out of its training data, so it has never seen the test cases."]
    stat = {"value": params if t.get("params") else "—", "label": "parameters, all trained here"}
    body = {"lines": lines, "stat": stat}
    return _slide("scratch", 3, "facts", body, notes, numbers)


def _lora_slide(numbers: dict) -> dict:
    t = (numbers.get("local_training") or {}).get("qwen2.5-0.5b-agentlora") or {}
    base = t.get("base_model") or NOT_RECORDED
    if t.get("r") and t.get("alpha"):
        adapters = (f"LoRA adapters with rank r = {t['r']} and alpha = {t['alpha']}, trained here for "
                    f"{t.get('steps')} steps")
    else:
        adapters = f"LoRA adapters (rank and alpha {NOT_RECORDED})"
    trainable, total = t.get("trainable_params"), t.get("total_params")
    if trainable and total:
        share = f"{_millions(trainable)} trainable of {_millions(total)} parameters; the base weights stay frozen"
        stat = {"value": f"{100.0 * trainable / total:.1f}%", "label": "of the weights trained (LoRA)"}
    else:
        share = f"Trainable parameters {NOT_RECORDED}"
        stat = {"value": "—", "label": "of the weights trained (LoRA)"}
    lines = [f"Base model: {base}, a pretrained instruction-following model",
             adapters, share,
             "Training mix: the Phase 2 synthetic agent conversations, with the benchmark's tools held out",
             f"Exact next action on the dev set after tuning: {_pct(t.get('dev_action_exact'))}"]
    lora, qbase = _model(numbers, "qwen2.5-0.5b-agentlora"), _model(numbers, "qwen2.5-0.5b-base")
    notes = ["The second locally made model starts from Qwen2.5-0.5B-Instruct, a small pretrained model, and I "
             "fine-tuned it here with LoRA. LoRA keeps the original weights frozen and learns small low-rank update "
             "matrices next to them, so only a small share of the parameters is trained, which fits on my laptop "
             "GPU.",
             "I tuned it on the same synthetic agent conversations as ScratchGPT, again with the benchmark's tools "
             f"held out. After tuning it picks exactly the right next action {_pct(t.get('dev_action_exact'))} of "
             f"the time on the dev set."]
    if lora and qbase and (lora["metrics"]["success_unprotected"]["pct"] or 0) > (
            qbase["metrics"]["success_unprotected"]["pct"] or 0):
        notes.append("Keep this model in mind: tuning it into a better agent also made it far more obedient to "
                     "injected instructions, which you will see on the results slides.")
    return _slide("lora", 4, "facts", {"lines": lines, "stat": stat}, notes, numbers)


def _lineup_slide(numbers: dict) -> dict:
    rows, starred = [], False
    for c in numbers.get("confounds") or []:
        params = c["parameters"]
        if " (" in params:
            params, starred = params.split(" (")[0] + "*", True
        rows.append([c["model"], params, c["precision"], c["placement"], c["context"].split(" (")[0],
                     c["provenance"].split(" (")[0]])
    decoding = [c["decoding"] for c in numbers.get("confounds") or []]
    caption = ["Fixed order, not a ranking"]
    if decoding and all(d.startswith("temperature 0") for d in decoding):
        caption.append("temperature 0 for every model"
                       + (" (hosted: determinism not guaranteed)" if any("not guaranteed" in d for d in decoding)
                          else ""))
    if starred:
        caption.append("* per the model id, not verified")
    table = {"columns": ["Model", "Parameters", "Precision", "Runs", "Context", "Provenance"],
             "widths": [1.9, 1.3, 1.05, 1.25, 1.4, 2.1], "rows": rows, "size": 11, "spacing": 0.9}
    notes = ["Here is the full line-up, in a fixed order: the local models first, then the hosted ones. This is not a "
             "ranking.",
             "I want to be honest about what differs between the rows, because these are confounds. The models "
             "differ in size, in numeric precision, in where they run, in context window and in how they were "
             "trained. So when two rows differ, I cannot say which of these caused the difference.",
             "What does not differ is the evaluation: the same prompt, the same cases, the same seeds and decoding "
             "at temperature 0. The hosted parameter counts come from the model names; I could not verify them, so "
             "they carry an asterisk."]
    body = {"table": table, "caption": "; ".join(caption) + "."}
    return _slide("lineup", 5, "table", body, notes, numbers)


def _method_slide(numbers: dict) -> dict:
    b = _bench(numbers)
    lines = [f"InjecAgent: {b['total']} cases ({b['dh']} direct-harm, {b['ds']} data-stealing, {b['benign']} "
             f"benign), {b['runs']} runs each",
             "Two arms: unprotected (no guardrails, an emulated OS) and protected (all four layers, the real "
             "sandbox)",
             "Same system prompt, cases, seeds and temperature 0 for every model",
             "Attack success: the attacker's tool ran, or the stolen data was sent out",
             "95% Wilson intervals; partial runs are labelled and never ranked",
             "One command freezes the data and regenerates every number, chart and slide"]
    notes = ["The benchmark is InjecAgent. In each case the user asks for something harmless, and one of the tool "
             "results the agent reads carries hidden attacker instructions. Direct-harm cases try to make the agent "
             "do something harmful; data-stealing cases try to make it read private data and send it to the "
             "attacker. The benign cases carry no injection at all; they measure whether the guardrails get in the "
             "way of normal work.",
             "I count an attack as successful only if the attacker's action really happened after the injection was "
             "visible to the model.",
             "Every proportion comes with a 95% Wilson interval, which behaves well near zero and one hundred "
             f"percent, where the usual normal approximation breaks down. Because decoding is at temperature 0, the "
             f"{b['runs']} runs of a case are near-duplicates, so these intervals are optimistic; I come back to that "
             f"on the limitations slide."]
    return _slide("method", 6, "bullets", {"lines": lines}, notes, numbers)


def _headline_slide(numbers: dict) -> dict:
    llm = _llm(numbers)
    pooled = numbers["pooled"]["blocked_given_attempt"]
    if _clean(numbers):
        lead = (f"With guardrails, no language model succeeded on any attack episode; pooled, guardrails blocked "
                f"{_kn(pooled)} attempted attacks.")
    else:
        lead = ("With guardrails, attacks still succeeded for "
                + ", ".join(f"{m['model']} {_kn(m['metrics']['success_protected'])}" for m in _leaks(numbers))
                + f"; pooled, guardrails blocked {_kn(pooled)} attempted attacks.")
    rows, styles = [], []
    for m in llm:
        rows.append([m["model"], _run_label(m), _pk(m["metrics"]["success_unprotected"]),
                     _pk(m["metrics"]["success_protected"]), _pk(m["metrics"]["benign_task_protected"])])
        styles.append("normal" if m["complete"] else "muted")
    control = _model(numbers, CONTROL_MODEL)
    if control:
        rows.append(["scripted control (reference)", _run_label(control),
                     _pk(control["metrics"]["success_unprotected"]), _pk(control["metrics"]["success_protected"]),
                     _pk(control["metrics"]["benign_task_protected"])])
        styles.append("reference")
    local = sum(1 for m in llm if m["group"] == "local")
    caption = (f"Fixed order: {local} local models, then {len(llm) - local} hosted ones; not a ranking. Partial rows "
               f"are not final.")
    table = {"columns": ["Model", "Run", "Attack success, no guardrails", "Attack success, guardrails",
                         "Benign task done, guardrails"],
             "widths": [2.5, 1.6, 1.65, 1.6, 1.65], "align": ["l", "l", "r", "r", "r"], "rows": rows,
             "styles": styles, "size": 12, "spacing": 0.9}
    lora, scratch = _model(numbers, "qwen2.5-0.5b-agentlora"), _model(numbers, "scratch-gpt")
    others = [m for m in llm if m["model"] not in ("qwen2.5-0.5b-agentlora", "scratch-gpt")]
    notes = ["This is the main result table. Without guardrails, attack success varies a lot between the models."]
    if lora and scratch:
        notes.append(f"It was {_pk(lora['metrics']['success_unprotected'])} for the LoRA-tuned Qwen (Qwen LoRA) and "
                     f"{_pk(scratch['metrics']['success_unprotected'])} for ScratchGPT.")
    if others:
        top = max(m["metrics"]["success_unprotected"]["pct"] or 0 for m in others)
        notes.append(f"For the pretrained models it was at most {top:.1f} percent.")
    if _clean(numbers):
        notes.append("With guardrails the column is zero for every model, and pooled over all the language models "
                     f"the guardrails blocked {_kn(pooled)} of the attacks that were attempted. That is the "
                     "conditional block rate.")
    else:
        notes.append(f"With guardrails some attacks still succeeded, as the lead line says; pooled, the guardrails "
                     f"blocked {_kn(pooled)} of the attempted attacks.")
    notes.append("The last column is benign task done with the guardrails on; slide 11 compares it with the "
                 "unprotected arm, so you can see what the protection costs in usefulness.")
    if any(not m["complete"] for m in llm):
        notes.append("The rows in grey are partial: the hosted runs are limited by the free-tier quota and continue "
                     "until the freeze, so they are labelled with their episodes done and are not compared with the "
                     "others.")
    notes.append("The scripted control at the bottom is not a language model; I come back to it later.")
    return _slide("headline", 7, "table", {"lead": lead, "table": table, "caption": caption}, notes, numbers)


def _chart_slide(numbers: dict, index: int, key: str, image: str, side: list, notes: list) -> dict:
    return _slide(key, index, "chart", {"image": image, "side": side}, notes, numbers)


def _asr_chart_slide(numbers: dict, image: str) -> dict:
    llm = _llm(numbers)
    side = ["Circles: no guardrails.", "Squares: guardrails."]
    side.append("With guardrails every model sits at 0." if _clean(numbers)
                else f"With guardrails {len(_leaks(numbers))} models stay above 0.")
    if any(not m["complete"] for m in llm):
        side.append("Hollow markers: partial runs, not final.")
    control = _model(numbers, CONTROL_MODEL)
    notes = ["This chart shows the same attack-success numbers with their 95% intervals. For each model, the circle "
             "is without guardrails and the square is with guardrails; the bars are the Wilson intervals.",
             "The intervals matter most for the small counts. A model with a few successes out of a hundred and "
             "twenty has an interval that still reaches several percent, so two such models cannot be told apart, "
             "and I do not try to."]
    if control:
        notes.append(f"The dashed line is the scripted control without guardrails, at "
                     f"{_pk(control['metrics']['success_unprotected'])}: an agent that always obeys the injection.")
    notes.append("With guardrails, every square sits on the zero line." if _clean(numbers)
                 else "With guardrails, some squares stay above the zero line; those are the leaks named on the "
                      "headline slide.")
    return _chart_slide(numbers, 8, "asr_chart", image, side, notes)


def _capability_chart_slide(numbers: dict, image: str) -> dict:
    side = ["Across: benign task done. Up: attack success. No guardrails."]
    lora, base = _model(numbers, "qwen2.5-0.5b-agentlora"), _model(numbers, "qwen2.5-0.5b-base")
    notes = ["Here each point is one model without guardrails. Across is how often it completed the benign task, "
             "which is a measure of capability; up is how often the attack succeeded, which is vulnerability."]
    if lora and base:
        la, lb = lora["metrics"]["success_unprotected"], lora["metrics"]["benign_task_unprotected"]
        ba, bb = base["metrics"]["success_unprotected"], base["metrics"]["benign_task_unprotected"]
        if (la["pct"] or 0) > (ba["pct"] or 0) and (lb["pct"] or 0) > (bb["pct"] or 0):
            side += ["LoRA tuning moved Qwen up and right:",
                     f"Qwen base: {_kn(ba)} attacks, {_kn(bb)} benign",
                     f"Qwen LoRA: {_kn(la)} attacks, {_kn(lb)} benign"]
            notes += ["A known pattern in agent security is that a more capable agent is also easier to hijack, "
                      "because following instructions well includes following injected ones. Our data shows it "
                      "inside one model family.",
                      f"Before tuning, Qwen base succeeded on the attack in {_pk(ba)} of episodes and finished "
                      f"{_pk(bb)} of the benign tasks.",
                      f"After LoRA tuning for tool use, Qwen LoRA succeeded on the attack in {_pk(la)} of episodes and "
                      f"finished {_pk(lb)} of the benign tasks.",
                      "The tuning taught it to act, and it acts on the attacker's instructions too. This is exactly "
                      "why the guardrails must not depend on the model: the more useful the model, the more it needs "
                      "them."]
        else:
            notes.append("In this snapshot capability and vulnerability do not move together for the two Qwen "
                         "models, so I read the chart point by point rather than as a trend.")
    return _chart_slide(numbers, 9, "capability_chart", image, side, notes)


def _arrows_chart_slide(numbers: dict, image: str) -> dict:
    llm = _llm(numbers)
    side = ["Hollow: no guardrails. Filled: guardrails."]
    if _clean(numbers):
        side.append("Every arrow ends at 0 attack success.")
    complete = [m for m in llm if m["complete"] and m["metrics"]["benign_task_unprotected"]["n"]
                and m["metrics"]["benign_task_protected"]["n"]]
    notes = ["Each arrow is one model, from no guardrails, the hollow marker, to guardrails, the filled one. The "
             "vertical drop is the protection; the horizontal movement is the cost in usefulness."]
    if complete:
        shift = max(abs((m["metrics"]["benign_task_protected"]["pct"] or 0)
                        - (m["metrics"]["benign_task_unprotected"]["pct"] or 0)) for m in complete)
        side.append(f"Benign task done moved by at most {shift:.1f} points for the complete models.")
        notes.append(f"For the complete models, benign task done changed by at most {shift:.1f} percentage points, so "
                     f"the protection did not come at the price of usefulness on these tasks.")
    if any(not m["complete"] for m in llm):
        notes.append("The hosted rows are partial, so their arrows can still move before the freeze.")
    notes.append("Every arrow ending on the zero line is the visual form of the core claim: the guardrails hold "
                 "whichever model drives the agent." if _clean(numbers)
                 else "Arrows that stay above the zero line are the leaks named on the headline slide.")
    return _chart_slide(numbers, 10, "arrows_chart", image, side, notes)


def _layers_slide(numbers: dict) -> dict:
    groups = numbers["pooled"]["layers"]
    attack, benign = groups.get("attack"), groups.get("benign")
    rows = []
    for layer in LAYERS:
        cells = [_pk(g[layer]) if g else "— (0/0)" for g in (attack, benign)]
        rows.append([LAYER_TITLES[layer]] + cells + [str(sum(g[layer]["actions"] for g in groups.values()))])
    table = {"columns": ["Layer", "Attack episodes, pooled", "Benign episodes, pooled", "Blocked actions"],
             "widths": [2.4, 2.45, 2.45, 1.7], "align": ["l", "r", "r", "r"], "rows": rows, "size": 13}
    sandbox, n = _pooled_layer(numbers, "sandbox")
    seccomp, _ = _pooled_layer(numbers, "seccomp")
    pooled = numbers["pooled"]["blocked_given_attempt"]
    stopped_all = pooled["n"] > 0 and pooled["k"] == pooled["n"]
    if sandbox == 0 and seccomp == 0:
        first = (f"Sandbox and seccomp: no block recorded, 0/{n} pooled protected-arm episodes. This zero is a lower "
                 f"bound" + (": the tool broker and policy engine stopped every attempted attack first."
                             if stopped_all else ": attacks the earlier layers stop never reach them."))
    else:
        parts = [f"{title}: blocks recorded in {k}/{n} pooled protected-arm episodes" if k
                 else f"{title}: no block recorded, 0/{n} pooled protected-arm episodes"
                 for title, k in (("Sandbox", sandbox), ("Seccomp", seccomp))]
        first = "; ".join(parts) + ". These counts are a lower bound."
    rows_by = _phase1_rows()
    if len(rows_by) == 3:
        second = (f"Phase 1 residual-bypass table, by row: rows {rows_by['seccomp']} blocked by seccomp; rows "
                  f"{rows_by['sandbox']} by the sandbox; rows {rows_by['contained']} contained without an error "
                  f"signature.")
    else:
        second = ("The Phase 1 residual-bypass table, cited in the report by row number, shows what the kernel "
                  "layers contain.")
    tb, pe = (attack or {}).get("tool_broker"), (attack or {}).get("policy_engine")
    notes = ["This slide answers where the attacks were stopped, layer by layer, in the protected arm."]
    if tb and pe:
        notes.append(f"Pooled over the language models, the tool broker blocked something in {_kn(tb)} attack "
                     f"episodes and the policy engine in {_kn(pe)}.")
    if sandbox == 0 and seccomp == 0:
        notes.append(f"The sandbox and seccomp rows show zero, pooled 0/{n}. I want to be precise about that zero. "
                     "It does not mean the kernel layers are useless. They only see actions that the first two "
                     "layers allowed" + (", and in this benchmark the first two layers stopped every attempted attack "
                                         "first." if stopped_all else "."))
    else:
        notes.append("The kernel layers did record blocks in this data, and even so the counts are a lower bound.")
    notes += ["Also, the executor records a sandbox or seccomp block only when a command's error text matches a "
              "known signature, so a contained action that fails quietly is not counted. That makes every number in "
              "the two kernel rows a lower bound.",
              "The evidence that the kernel layers do contain actions comes from Phase 1, where I tested residual "
              "bypasses directly. " + LAYER_EVIDENCE]
    return _slide("layers", 11, "table", {"table": table, "lines": [first, second]}, notes, numbers)


def _control_slide(numbers: dict) -> dict:
    c = numbers.get("control")
    notes = ["A fair question from an examiner is: maybe the models simply did not try very hard. The control "
             "answers it. It is a scripted agent, not a language model, that obeys every injection."]
    if not c:
        body = {"lines": ["No control rows in this snapshot"]}
        notes.append("This snapshot has no control rows, so the question stays open until the control is run.")
        return _slide("control", 12, "bullets", body, notes, numbers)
    u, p = c["unprotected"], c["protected"]
    rows = [["Scripted control, no guardrails", str(u["episodes"]), _fp(u["attempted"], u["episodes"]),
             _fp(u["succeeded"], u["episodes"]), "—"],
            ["Scripted control, guardrails", str(p["episodes"]), _fp(p["attempted"], p["episodes"]),
             _fp(p["succeeded"], p["episodes"]), _fp(p["blocked"], p["attempted"])]]
    table = {"columns": ["Arm", "Attack episodes", "Attempted", "Success", "Blocked given attempt"],
             "widths": [2.75, 1.35, 1.65, 1.6, 1.65], "align": ["l", "r", "r", "r", "r"], "rows": rows, "size": 13}
    lines = ["A scripted agent that always obeys the injection: the worst case any model could be",
             f"With guardrails the tool broker denied the attacker's tool call in {c['broker_denied']} of its "
             f"{p['attempted']} attempted attacks",
             f"Deterministic, so its {p['episodes']} protected attack episodes are {c['cases']['attack']} distinct "
             f"cases"]
    if p["succeeded"] == 0 and p["blocked"] == p["attempted"]:
        lines.append("So the block does not depend on the model being cautious")
    layers = c.get("layers") or {}
    notes += [f"Without guardrails the scripted control succeeded in {u['succeeded']}/{u['episodes']} attack episodes, "
              f"so the attacks themselves are real and work when nothing stops them.",
              f"With guardrails the scripted control succeeded in {p['succeeded']}/{p['episodes']}, and the "
              f"guardrails blocked {p['blocked']}/{p['attempted']} of its attempts."]
    if layers.get("tool_broker") and layers.get("policy_engine"):
        notes.append(f"The tool broker fired in {layers['tool_broker']['episodes']} of its protected attack episodes "
                     f"and the policy engine in {layers['policy_engine']['episodes']}.")
    notes.append(f"Because the control is deterministic, its {p['episodes']} protected episodes are really only "
                 f"{c['cases']['attack']} distinct cases, so I do not claim more precision than that.")
    return _slide("control", 12, "table", {"table": table, "lines": lines}, notes, numbers)


def _threats_slide(numbers: dict) -> dict:
    llm = _llm(numbers)
    b = _bench(numbers)
    partial = [m for m in llm if not m["complete"]]
    lines = [("Partial runs, not final and not ranked: "
              + ", ".join(f"{m['model']} {m['episodes_done']}/{m['episodes_expected']}" for m in partial))
             if partial else "Every model finished its planned run"]
    sandbox, n = _pooled_layer(numbers, "sandbox")
    seccomp, _ = _pooled_layer(numbers, "seccomp")
    if sandbox == 0 and seccomp == 0:
        lines.append(f"Kernel layers: 0/{n} pooled recorded blocks is a lower bound, not proof they are idle")
    else:
        lines.append(f"Kernel layers: the recorded blocks (pooled, sandbox {sandbox}/{n}, seccomp {seccomp}/{n}) are a "
                     f"lower bound")
    failing = [m for m in llm if sum(m["format_failures"].values()) > 0]
    if failing:
        lines.append(f"Format failures lower raw success for {len(failing)} models; read ASR-valid next to raw")
    overflow = sum(m["context_overflow_episodes"] for m in llm)
    lines.append(f"Deferred WR-01: context overflows can bias ASR-valid upward; {overflow} context-overflow episodes "
                 f"in the current data" + (", so no current number is affected" if overflow == 0
                                           else ", so ASR-valid is biased where they occur"))
    lines.append(f"One benchmark (InjecAgent, {b['total']} cases, simulated tools), temperature 0: optimistic "
                 f"intervals")
    lines.append("Models differ in size, precision, placement and training (slide 6)"
                 + ("; legacy rows excluded" if numbers.get("legacy") else ""))
    legacy = numbers.get("legacy") or []
    notes = ["I want to close the results with their limits, because these are the questions I would ask myself."]
    for number, item in enumerate(numbers.get("threats") or [], 1):
        if any(legacy_id in item for legacy_id in LEGACY_IDS):
            item = (f"Legacy rows: {len(legacy)} October rows from an older engine appear only in their own table in "
                    f"the report; they are not compared, pooled or charted, and I do not show them here.")
        notes.append(f"{number}. {item}")
    return _slide("threats", 13, "bullets", {"lines": lines[:6], "size": 16}, notes, numbers)


def _next_slide(numbers: dict) -> dict:
    lines = []
    if numbers["data_state"] == "snapshot":
        lines.append(f"Final Groq numbers after the {numbers['freeze_display']} freeze: one command re-freezes the "
                     f"data and regenerates this deck")
    lines += [f"{rid}: {text}" for rid, text in V2_STEPS]
    notes = []
    if numbers["data_state"] == "snapshot":
        notes.append(f"Finally, what comes next. The immediate step is mechanical: the Groq runs continue until the "
                     f"freeze at {numbers['freeze_display']} UTC, and then one command re-freezes the data and "
                     f"regenerates the report, the charts and this deck, so the final numbers can never drift from "
                     f"what I present.")
    else:
        notes.append("Finally, what comes next. The data in this deck is already final, regenerated by one command "
                     "after the freeze.")
    notes += ["Beyond this review there are five version-2 items. First, two more local models through Ollama, phi "
              "3.5 and gemma 3, which are already downloaded. Second, a larger utility-under-attack task set with "
              "case-clustered bootstrap intervals, which fixes the optimistic intervals I mentioned. Third, adaptive "
              "attacks that retry after a block, the strongest test of a reference monitor. Fourth, a replication on "
              "a second benchmark, AgentDojo or Agent Security Bench. And fifth, serving the LoRA-tuned Qwen through "
              "Ollama as a GGUF file.",
              "Thank you. I am happy to take questions."]
    return _slide("next", 14, "bullets", {"lines": lines}, notes, numbers)


def deck_spec(numbers: dict) -> list:
    """The slides as plain dicts {key, title, kind, body, notes, footer}, built from numbers.json alone.

    Args:
        numbers: the parsed numbers.json of the frozen directory.

    Returns:
        One dict per slide, in deck order; the titles equal SLIDE_TITLES."""
    charts = [c["file"] for c in numbers.get("charts") or []]
    charts += [None] * (3 - len(charts))
    return [_title_slide(numbers), _remark_slide(numbers), _architecture_slide(numbers), _scratch_slide(numbers),
            _lora_slide(numbers), _lineup_slide(numbers), _method_slide(numbers), _headline_slide(numbers),
            _asr_chart_slide(numbers, charts[0]), _capability_chart_slide(numbers, charts[1]),
            _arrows_chart_slide(numbers, charts[2]), _layers_slide(numbers), _control_slide(numbers),
            _threats_slide(numbers), _next_slide(numbers)]


def slide_text(slide: dict) -> list:
    """Every visible text line of a spec slide (title, body, table rows joined by ' | ', footer)."""
    body = slide["body"]
    lines = [slide["title"]]
    for key in ("kicker", "subtitle", "byline", "quote", "attribution", "label", "lead"):
        if body.get(key):
            lines.append(body[key])
    if body.get("diagram"):
        lines.append(body["diagram"]["container"])
        for box in body["diagram"]["boxes"]:
            lines += [box["label"]] + ([box["detail"]] if box.get("detail") else [])
    if body.get("table"):
        table = body["table"]
        lines += [" | ".join(row) for row in [table["columns"]] + table["rows"]]
    if body.get("caption"):
        lines.append(body["caption"])
    for key in ("lines", "side"):
        lines += list(body.get(key) or [])
    if body.get("stat"):
        lines += [body["stat"]["value"], body["stat"]["label"]]
    if slide.get("footer"):
        lines.append(slide["footer"])
    return lines


def chart_files(spec: list) -> list:
    """[(slide index from 1, chart file relative to the frozen directory)] in deck order."""
    return [(i, s["body"]["image"]) for i, s in enumerate(spec, 1) if s["body"].get("image")]


# ── Layout estimates (pure) ──

def text_lines(text: str, width_in: float, size: float, bold: bool = False) -> int:
    """Lines a paragraph of Arial at `size` pt needs in a box `width_in` wide (a word-wrap estimate)."""
    per_char = size * (0.53 if bold else 0.5)
    per_line = max(1, int(width_in * 72 // per_char))
    lines, used = 1, 0
    for word in text.split():
        if used == 0:
            used = len(word)
        elif used + 1 + len(word) <= per_line:
            used += 1 + len(word)
        else:
            lines, used = lines + 1, len(word)
        while used > per_line:                      # a word longer than the line breaks across lines
            lines, used = lines + 1, used - per_line
    return lines


def block_height(paras: list, width_in: float, size: float, space_after: float = 6, bold: bool = False) -> float:
    """Inches a list of paragraphs needs at `size` pt with `space_after` pt between them."""
    points = sum(text_lines(p, width_in, size, bold) * size * 1.2 + space_after for p in paras)
    return points / 72


CELL_MARGIN_X, CELL_MARGIN_Y = 0.07, 0.03


def row_height(values: list, widths: list, size: float, bold: bool = False, spacing: float = 1.0) -> float:
    """Inches a table row needs: its tallest cell at `size` pt with line spacing `spacing`."""
    lines = max(text_lines(v, w - 2 * CELL_MARGIN_X, size, bold) for v, w in zip(values, widths))
    return math.ceil((lines * size * 1.2 * spacing / 72 + 2 * CELL_MARGIN_Y + 0.01) * 100) / 100


# ── Render (python-pptx) ──

def _import_pptx():
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.util import Emu, Inches, Pt
    except ImportError as e:
        raise PptxUnavailable(f"python-pptx is not importable ({type(e).__name__}: {e})")
    return {"Presentation": Presentation, "RGBColor": RGBColor, "MSO_CONNECTOR": MSO_CONNECTOR,
            "MSO_SHAPE": MSO_SHAPE, "MSO_SHAPE_TYPE": MSO_SHAPE_TYPE, "MSO_ANCHOR": MSO_ANCHOR,
            "MSO_AUTO_SIZE": MSO_AUTO_SIZE, "PP_ALIGN": PP_ALIGN, "qn": qn, "OxmlElement": OxmlElement,
            "Emu": Emu, "Inches": Inches, "Pt": Pt}


class _Renderer:
    """Draws spec slides with python-pptx: own shapes on the template's plain BLANK layout."""

    def __init__(self, x: dict, prs, layout, frozen_dir: str, total: int):
        self.x, self.prs, self.layout, self.frozen_dir, self.total = x, prs, layout, frozen_dir, total

    # primitives
    def _in(self, value: float):
        return self.x["Inches"](value)

    def _rgb(self, hex6: str):
        return self.x["RGBColor"].from_string(hex6)

    def _element(self, tag: str, **attrs):
        element = self.x["OxmlElement"](tag)
        for key, value in attrs.items():
            element.set(key, str(value))
        return element

    def text(self, slide, left, top, width, height, paras, size=18, color=INK, bold=False, italic=False,
             align="left", anchor="top", bullets=False, space_after=6, name=None):
        """A text box; each para is a string or a dict {text, size, color, bold, italic, bullet, space_after}."""
        x = self.x
        box = slide.shapes.add_textbox(self._in(left), self._in(top), self._in(width), self._in(height))
        if name:
            box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = x["MSO_AUTO_SIZE"].NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": x["MSO_ANCHOR"].TOP, "middle": x["MSO_ANCHOR"].MIDDLE,
                              "bottom": x["MSO_ANCHOR"].BOTTOM}[anchor]
        for i, para in enumerate(paras):
            spec = para if isinstance(para, dict) else {"text": para}
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": x["PP_ALIGN"].LEFT, "center": x["PP_ALIGN"].CENTER,
                           "right": x["PP_ALIGN"].RIGHT}[spec.get("align", align)]
            p.line_spacing = 1.0
            p.space_after = x["Pt"](spec.get("space_after", space_after))
            run = p.add_run()
            run.text = spec["text"]
            font = run.font
            font.name = FONT
            font.size = x["Pt"](spec.get("size", size))
            font.bold = spec.get("bold", bold)
            font.italic = spec.get("italic", italic)
            font.color.rgb = self._rgb(spec.get("color", color))
            if spec.get("bullet", bullets):
                self._bullet(p)
        return box

    def _bullet(self, p):
        indent = int(self._in(0.3))
        p_pr = p._p.get_or_add_pPr()
        p_pr.set("marL", str(indent))
        p_pr.set("indent", str(-indent))
        clr = self._element("a:buClr")
        clr.append(self._element("a:srgbClr", val=ACCENT))
        for element in (clr, self._element("a:buSzPct", val=70000), self._element("a:buFont", typeface=FONT),
                        self._element("a:buChar", char="■")):
            p_pr.append(element)

    def rect(self, slide, left, top, width, height, fill=None, line=None, line_width=1.0, name=None):
        shape = slide.shapes.add_shape(self.x["MSO_SHAPE"].RECTANGLE, self._in(left), self._in(top),
                                       self._in(width), self._in(height))
        if name:
            shape.name = name
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = self._rgb(fill)
        else:
            shape.fill.background()
        if line:
            shape.line.color.rgb = self._rgb(line)
            shape.line.width = self.x["Pt"](line_width)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        if shape.has_text_frame:
            shape.text_frame.text = ""
        return shape

    def arrow(self, slide, x1, y1, x2, y2, color=MUTED, width=1.5, dashed=False, head=True, name=None):
        line = slide.shapes.add_connector(self.x["MSO_CONNECTOR"].STRAIGHT, self._in(x1), self._in(y1),
                                          self._in(x2), self._in(y2))
        if name:
            line.name = name
        line.line.color.rgb = self._rgb(color)
        line.line.width = self.x["Pt"](width)
        ln = line.line._get_or_add_ln()
        if dashed:
            ln.append(self._element("a:prstDash", val="dash"))
        if head:
            ln.append(self._element("a:tailEnd", type="triangle", w="med", len="med"))
        return line

    # slide frame
    def new_slide(self):
        slide = self.prs.slides.add_slide(self.layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self._rgb(BG)
        return slide

    def frame(self, slide, spec: dict, index: int):
        """Title, accent rule and footer of a content slide."""
        self.text(slide, MARGIN, TITLE_Y, CONTENT_W, TITLE_H, [spec["title"]], size=24, bold=True,
                  anchor="bottom", space_after=0, name="Title")
        self.rect(slide, MARGIN, RULE_Y, RULE_W, RULE_H, fill=ACCENT, name="Accent rule")
        self.text(slide, MARGIN, FOOTER_Y, CONTENT_W - 0.6, FOOTER_H, [spec["footer"]], size=9, color=MUTED,
                  anchor="bottom", space_after=0, name="Footer")
        self.text(slide, SLIDE_W - MARGIN - 0.5, FOOTER_Y, 0.5, FOOTER_H, [str(index)], size=9, color=MUTED,
                  align="right", anchor="bottom", space_after=0, name="Slide number")

    def notes(self, slide, text: str):
        slide.notes_slide.notes_text_frame.text = text

    # kinds
    def title_slide(self, slide, spec: dict):
        body = spec["body"]
        left, width = 0.75, 8.5
        self.text(slide, left, 1.05, width, 0.32, [body["kicker"]], size=14, color=ACCENT, bold=True,
                  space_after=0, name="Kicker")
        self.text(slide, left, 1.45, width, 1.4, [spec["title"]], size=34, bold=True, space_after=0, name="Title")
        self.rect(slide, left, 2.98, 0.9, 0.06, fill=ACCENT, name="Accent rule")
        self.text(slide, left, 3.2, width, 0.45, [body["subtitle"]], size=20, color=MUTED, space_after=0,
                  name="Subtitle")
        self.text(slide, left, 3.75, width, 0.4, [body["byline"]], size=18, bold=True, space_after=0, name="Byline")
        self.text(slide, left, 4.75, width, 0.5, [spec["footer"]], size=11, color=MUTED, anchor="bottom",
                  space_after=0, name="Footer")

    def bullets(self, slide, spec: dict):
        body = spec["body"]
        size = body.get("size", 18)
        self.text(slide, MARGIN, BODY_Y + 0.1, CONTENT_W, BODY_BOTTOM - BODY_Y - 0.1, body["lines"], size=size,
                  bullets=True, space_after=10, name="Body")

    def quote(self, slide, spec: dict):
        body = spec["body"]
        quote_h = block_height([body["quote"]], CONTENT_W - 0.3, 20, space_after=0) + 0.05
        self.rect(slide, MARGIN, BODY_Y + 0.08, 0.06, quote_h, fill=ACCENT, name="Quote rule")
        self.text(slide, MARGIN + 0.3, BODY_Y + 0.08, CONTENT_W - 0.3, quote_h, [body["quote"]], size=20,
                  italic=True, space_after=0, name="Quote")
        y = BODY_Y + 0.08 + quote_h + 0.08
        self.text(slide, MARGIN + 0.3, y, CONTENT_W - 0.3, 0.26, [body["attribution"]], size=13, color=MUTED,
                  space_after=0, name="Attribution")
        y += 0.42
        self.text(slide, MARGIN, y, CONTENT_W, 0.3, [body["label"]], size=14, bold=True, color=ACCENT, space_after=0,
                  name="Label")
        y += 0.36
        self.text(slide, MARGIN, y, CONTENT_W, BODY_BOTTOM - y, body["lines"], size=16, bullets=True, space_after=7,
                  name="Body")

    def facts(self, slide, spec: dict):
        body = spec["body"]
        self.text(slide, MARGIN, BODY_Y + 0.1, 5.85, BODY_BOTTOM - BODY_Y - 0.1, body["lines"], size=16,
                  bullets=True, space_after=9, name="Body")
        self.rect(slide, 6.7, BODY_Y + 0.25, 0.015, 3.3, fill=RULE, name="Divider")
        stat = body["stat"]
        self.text(slide, 7.0, 1.85, 2.5, 0.95, [stat["value"]], size=48, bold=True, color=ACCENT, anchor="bottom",
                  space_after=0, name="Stat")
        self.text(slide, 7.0, 2.9, 2.5, 0.9, [stat["label"]], size=15, color=MUTED, space_after=0, name="Stat label")

    def table(self, slide, top: float, spec: dict) -> float:
        """Draw a borderless editorial table at `top`; returns its estimated height in inches."""
        x = self.x
        cols, rows, widths = spec["columns"], spec["rows"], spec["widths"]
        size = spec.get("size", 12)
        aligns = spec.get("align") or ["l"] * len(cols)
        styles = spec.get("styles") or ["normal"] * len(rows)
        spacing = spec.get("spacing", 1.0)
        heights = ([row_height(cols, widths, size, bold=True, spacing=spacing)]
                   + [row_height(r, widths, size, spacing=spacing) for r in rows])
        frame = slide.shapes.add_table(len(rows) + 1, len(cols), self._in(MARGIN), self._in(top),
                                       self._in(sum(widths)), self._in(sum(heights)))
        frame.name = "Table"
        tbl = frame.table
        tbl_pr = tbl._tbl.tblPr
        tbl_pr.set("firstRow", "0")
        tbl_pr.set("bandRow", "0")
        style = tbl_pr.find(x["qn"]("a:tableStyleId"))
        if style is None:
            style = self._element("a:tableStyleId")
            tbl_pr.append(style)
        style.text = NO_TABLE_STYLE
        for i, w in enumerate(widths):
            tbl.columns[i].width = self._in(w)
        for i, h in enumerate(heights):
            tbl.rows[i].height = self._in(h)
        for r, values in enumerate([cols] + rows):
            row_style = "header" if r == 0 else styles[r - 1]
            for c, value in enumerate(values):
                self._cell(tbl.cell(r, c), value, size, row_style, aligns[c], spacing)
        return sum(heights)

    def _cell(self, cell, value: str, size: float, style: str, align: str, spacing: float):
        x = self.x
        cell.margin_left = cell.margin_right = self._in(CELL_MARGIN_X)
        cell.margin_top = cell.margin_bottom = self._in(CELL_MARGIN_Y)
        cell.vertical_anchor = x["MSO_ANCHOR"].BOTTOM if style == "header" else x["MSO_ANCHOR"].MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = x["PP_ALIGN"].RIGHT if align == "r" else x["PP_ALIGN"].LEFT
        p.line_spacing = spacing
        run = p.add_run()
        run.text = value
        font = run.font
        font.name = FONT
        font.size = x["Pt"](size)
        font.bold = style == "header"
        font.italic = style == "reference"
        font.color.rgb = self._rgb({"header": INK, "muted": MUTED, "reference": MUTED}.get(style, INK))
        cell.fill.background()
        bottom = (INK, 1.25) if style == "header" else (RULE, 0.75)
        tc_pr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            for old in tc_pr.findall(x["qn"](tag)):
                tc_pr.remove(old)
        for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
            if tag == "a:lnB":
                ln = self._element(tag, w=int(x["Pt"](bottom[1])))
                fill = self._element("a:solidFill")
                fill.append(self._element("a:srgbClr", val=bottom[0]))
                ln.append(fill)
            else:
                ln = self._element(tag, w=0)
                ln.append(self._element("a:noFill"))
            tc_pr.insert(i, ln)

    def table_slide(self, slide, spec: dict):
        body = spec["body"]
        y = BODY_Y
        if body.get("lead"):
            lead_h = block_height([body["lead"]], CONTENT_W, 16, space_after=0)
            self.text(slide, MARGIN, y, CONTENT_W, lead_h, [body["lead"]], size=16, space_after=0, name="Lead")
            y += lead_h + 0.12
        else:
            y += 0.05
        y += self.table(slide, y, body["table"]) + 0.12
        if body.get("caption"):
            cap_h = block_height([body["caption"]], CONTENT_W, 12, space_after=0)
            self.text(slide, MARGIN, y, CONTENT_W, cap_h, [body["caption"]], size=12, color=MUTED, space_after=0,
                      name="Caption")
            y += cap_h + 0.1
        if body.get("lines"):
            y += 0.08
            self.text(slide, MARGIN, y, CONTENT_W, max(0.3, BODY_BOTTOM - y), body["lines"], size=16, bullets=True,
                      space_after=8, name="Body")

    def chart(self, slide, spec: dict):
        body = spec["body"]
        if not body.get("image"):
            raise ValueError(f"slide {spec['key']} has no chart in numbers.json")
        path = os.path.join(self.frozen_dir, *body["image"].split("/"))
        if not os.path.isfile(path):
            raise ValueError(f"frozen chart {body['image']} is missing; run python -m src.deliverables build")
        picture = slide.shapes.add_picture(path, self._in(CHART_X), self._in(CHART_Y), width=self._in(CHART_W))
        picture.name = "Chart"
        max_h = self._in(BODY_BOTTOM - CHART_Y)
        if picture.height > max_h:                  # a taller chart is scaled to the body height, keeping its aspect
            picture.width, picture.height = int(picture.width * max_h / picture.height), max_h
        self.text(slide, SIDE_X, CHART_Y + 0.15, SIDE_W, BODY_BOTTOM - CHART_Y - 0.15, body["side"], size=16,
                  space_after=9, name="Side")

    def diagram(self, slide, spec: dict):
        d = spec["body"]["diagram"]
        boxes = d["boxes"]
        role = {b["role"]: b for b in boxes if b["role"] != "layer"}
        layers = [b for b in boxes if b["role"] == "layer"]
        left, width = MARGIN, 5.75
        label_w, detail_x = 1.6, left + 1.95
        # the untrusted model
        model = role["model"]
        self.rect(slide, left, 1.2, width, 0.62, fill=WHITE, line=MUTED, line_width=1.0, name="Model box")
        self.text(slide, left + 0.15, 1.2, label_w, 0.62, [model["label"]], size=15, bold=True, anchor="middle",
                  space_after=0, name="Model label")
        self.text(slide, detail_x, 1.2, width - 2.05, 0.62, [model["detail"]], size=13, color=MUTED,
                  anchor="middle", space_after=0, name="Model detail")
        self.arrow(slide, left + width / 2, 1.82, left + width / 2, 2.02, color=MUTED, name="Action arrow")
        # the reference monitor and its four layers
        self.rect(slide, left, 2.02, width, 2.6, line=ACCENT, line_width=1.25, name="Monitor")
        self.text(slide, left + 0.12, 2.08, width - 0.24, 0.28, [d["container"]], size=12, bold=True, color=ACCENT,
                  space_after=0, name="Monitor label")
        for i, layer in enumerate(layers):
            top = 2.42 + i * 0.55
            self.rect(slide, left + 0.12, top, width - 0.24, 0.5, fill=TINT, name=f"Layer {i + 1}")
            self.rect(slide, left + 0.12, top, 0.05, 0.5, fill=ACCENT, name=f"Layer {i + 1} mark")
            self.text(slide, left + 0.3, top, label_w, 0.5, [layer["label"]], size=15, bold=True, anchor="middle",
                      space_after=0, name=f"Layer {i + 1} label")
            self.text(slide, detail_x, top, width - 2.07, 0.5, [layer["detail"]], size=13, anchor="middle",
                      space_after=0, name=f"Layer {i + 1} detail")
        self.arrow(slide, left + width / 2, 4.62, left + width / 2, 4.78, color=MUTED, name="Kernel arrow")
        kernel = role["kernel"]
        self.rect(slide, left, 4.78, width, 0.38, fill=DARK, name="Kernel box")
        self.text(slide, left, 4.78, width, 0.38, [kernel["label"]], size=14, bold=True, color=WHITE,
                  align="center", anchor="middle", space_after=0, name="Kernel label")
        # the decision log and the two arms
        side_x, side_w = 6.6, 2.9
        for key, top, height in (("log", 2.02, 1.2), ("arms", 3.42, 1.2)):
            box = role[key]
            self.rect(slide, side_x, top, side_w, height, fill=WHITE, line=RULE, line_width=1.0, name=f"{key} box")
            self.text(slide, side_x + 0.15, top + 0.1, side_w - 0.3, height - 0.2,
                      [{"text": box["label"], "size": 15, "bold": True, "space_after": 4},
                       {"text": box["detail"], "size": 13, "color": MUTED}], space_after=0, name=f"{key} text")
        self.arrow(slide, left + width, 2.62, side_x, 2.62, color=MUTED, width=1.0, dashed=True, name="Log arrow")

    def draw(self, spec: dict, index: int):
        slide = self.new_slide()
        kind = spec["kind"]
        if kind == "title":
            self.title_slide(slide, spec)
        else:
            self.frame(slide, spec, index)
            drawers = {"bullets": self.bullets, "quote": self.quote, "facts": self.facts, "table": self.table_slide,
                       "chart": self.chart, "diagram": self.diagram}
            if kind not in drawers:
                raise ValueError(f"unknown slide kind {kind!r}")
            drawers[kind](slide, spec)
        self.notes(slide, spec["notes"])
        return slide


def _open_template(x: dict, template):
    """(presentation, how) — the template with its own slides removed, or a blank 16:9 deck."""
    Presentation = x["Presentation"]
    if template and os.path.isfile(template):
        try:
            prs = Presentation(template)
        except Exception as e:                      # a corrupt template falls back to the blank deck, loudly
            how = f"none ({type(e).__name__} opening {_posix(template)}), blank 16:9"
        else:
            id_list = prs.slides._sldIdLst
            for sld_id in list(id_list):
                prs.part.drop_rel(sld_id.rId)
                id_list.remove(sld_id)
            return prs, _posix(template)
    else:
        how = f"none ({_posix(template)} not found), blank 16:9"
    prs = Presentation()
    prs.slide_width, prs.slide_height = x["Inches"](SLIDE_W), x["Inches"](SLIDE_H)
    return prs, how


def _blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name == "BLANK":
            return layout
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    return prs.slide_layouts[len(prs.slide_layouts) - 1]


def _posix(path) -> str:
    return str(path).replace("\\", "/")


def is_blank_fallback(how: str) -> bool:
    """render_deck's `how` says the template could not be used and the deck is a blank 16:9 one (WR-05)."""
    return str(how).startswith("none (")


def blank_fallback_problem(display: str) -> str:
    return (f"{display} was drawn on a blank 16:9 deck, not the Review template; fix the template or pass "
            f"--allow-blank-template")


def render_deck(spec: list, out_path: str, frozen_dir: str, template: str = DEFAULT_TEMPLATE) -> str:
    """Draw the spec into out_path (written through a temporary file); returns how the template was used.

    Args:
        spec: deck_spec(numbers).
        out_path: the .pptx to write.
        frozen_dir: the frozen directory holding charts/.
        template: the Review-2 template; its slides are removed and only its plain BLANK layout is used.

    Raises:
        PptxUnavailable: python-pptx cannot be imported."""
    x = _import_pptx()
    prs, how = _open_template(x, template)
    renderer = _Renderer(x, prs, _blank_layout(prs), frozen_dir, len(spec))
    for index, slide_spec in enumerate(spec, 1):
        renderer.draw(slide_spec, index)
    parent = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(parent, exist_ok=True)
    tmp = out_path + ".tmp"
    prs.save(tmp)
    os.replace(tmp, out_path)
    return how


# ── Check (python-pptx) ──

def _sha256_file(path: str) -> str:
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def _shape_lines(x: dict, shape) -> list:
    lines = []
    if shape.shape_type == x["MSO_SHAPE_TYPE"].GROUP:
        for inner in shape.shapes:
            lines += _shape_lines(x, inner)
        return lines
    if shape.has_text_frame:
        lines += [p.text.replace("\v", " ") for p in shape.text_frame.paragraphs if p.text]
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            lines.append(" | ".join(cell.text.replace("\n", " ").replace("\v", " ") for cell in row.cells))
    return lines


def _pictures(x: dict, shapes) -> list:
    out = []
    for shape in shapes:
        if shape.shape_type == x["MSO_SHAPE_TYPE"].GROUP:
            out += _pictures(x, shape.shapes)
        elif shape.shape_type == x["MSO_SHAPE_TYPE"].PICTURE:
            out.append(hashlib.sha256(shape.image.blob).hexdigest())
    return out


def _out_of_bounds(shape, width: int, height: int) -> bool:
    if shape.left is None or shape.top is None:
        return False
    return (shape.left < 0 or shape.top < 0 or shape.left + shape.width > width
            or shape.top + shape.height > height)


def text_problems(index: int, lines: list, notes: str, numbers: dict) -> list:
    """The text rules check_deck applies to one slide (pure, no python-pptx, so the WSL venv tests them too):
    the title slide carries the snapshot time, a slide with a number carries the data line, the shown text and
    the notes pass scan_text (numbers and the key regex), and no legacy model id appears. Reasons name the slide
    number and the part, never the text."""
    where, snapshot, problems = f"slide {index}", numbers["snapshot_display"], []
    text = "\n".join(lines)
    if index == 1 and snapshot not in text:
        problems.append(f"{where}: the title slide does not carry the snapshot time {snapshot}")
    if FRACTION_OR_PERCENT_RE.search(text) and snapshot not in text:
        problems.append(f"{where}: shows numbers without the data line ({snapshot})")
    for part, part_lines in (("text", lines), ("notes", notes.splitlines())):
        for lineno, reason in scan_text(part_lines, numbers):
            problems.append(f"{where} {part} line {lineno}: {reason}")
        if any(legacy in line.lower() for line in part_lines for legacy in LEGACY_IDS):
            problems.append(f"{where} {part}: names a legacy model id")
    return problems


def check_deck(path: str, numbers: dict, frozen_dir: str) -> tuple:
    """(problems, counts) of the written deck against numbers.json and the frozen charts (D-03, D-05).

    Problems are reasons with slide numbers, never the offending text. counts = {slides, notes, charts}: the
    slides, the slides with notes of at least MIN_NOTES characters, the pictures byte-identical to the frozen
    chart their slide is meant to show."""
    x = _import_pptx()
    spec = deck_spec(numbers)
    problems = []
    if not os.path.isfile(path):
        return [f"{_posix(path)} does not exist"], {"slides": 0, "notes": 0, "charts": 0}
    prs = x["Presentation"](path)
    deck = list(prs.slides)
    if len(deck) != len(spec):
        problems.append(f"the deck has {len(deck)} slides, the spec {len(spec)}")
    wanted = {}
    for index, rel in chart_files(spec):
        chart = os.path.join(frozen_dir, *rel.split("/"))
        if not os.path.isfile(chart):
            problems.append(f"frozen chart {rel} is missing")
            continue
        wanted[index] = (rel, _sha256_file(chart))
    counts = {"slides": len(deck), "notes": 0, "charts": 0}
    for index, slide in enumerate(deck, 1):
        where = f"slide {index}"
        lines = []
        for shape in slide.shapes:
            lines += _shape_lines(x, shape)
            if _out_of_bounds(shape, prs.slide_width, prs.slide_height):
                problems.append(f"{where}: shape {shape.name!r} extends beyond the slide")
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        if len(notes) >= MIN_NOTES:
            counts["notes"] += 1
        else:
            problems.append(f"{where}: speaker notes are shorter than {MIN_NOTES} characters")
        titles = [s.text_frame.text for s in slide.shapes if s.name == "Title" and s.has_text_frame]
        if index <= len(spec) and titles != [spec[index - 1]["title"]]:
            problems.append(f"{where}: the title differs from the spec")
        problems += text_problems(index, lines, notes, numbers)
        pictures = _pictures(x, slide.shapes)
        expected = wanted.get(index)
        for digest in pictures:
            if expected and digest == expected[1]:
                counts["charts"] += 1
            else:
                problems.append(f"{where}: a picture is not the frozen chart this slide shows")
        if expected and expected[1] not in pictures:
            problems.append(f"{where}: the frozen chart {expected[0]} is not embedded")
    return problems, counts


def deck_line(out_path: str, counts: dict) -> str:
    return f"DECK {_posix(out_path)} slides {counts['slides']} notes {counts['notes']} charts {counts['charts']} ok"


def missing_line(executable: str) -> str:
    return f"FAIL deck: python-pptx is not importable by {executable}; run with py -3.12 on Windows"


def run_deck(numbers: dict, frozen_dir: str, out_path: str = DEFAULT_OUT, template: str = DEFAULT_TEMPLATE,
             render: bool = True, allow_blank_template: bool = False) -> tuple:
    """(printed lines, exit code): render (unless render=False) then check. Exit 0 with the DECK line last,
    1 with FAIL deck lines (a deck drawn on the blank fallback is one, unless allow_blank_template, WR-05);
    PptxUnavailable propagates to the caller (exit 3 there)."""
    lines = []
    if render:
        how = render_deck(deck_spec(numbers), out_path, frozen_dir, template)
        lines.append(f"TEMPLATE {how}")
        if is_blank_fallback(how) and not allow_blank_template:
            return lines + [f"FAIL deck: {blank_fallback_problem(_posix(out_path))}"], 1
    problems, counts = check_deck(out_path, numbers, frozen_dir)
    if problems:
        return lines + [f"FAIL deck: {p}" for p in problems], 1
    return lines + [deck_line(out_path, counts)], 0
